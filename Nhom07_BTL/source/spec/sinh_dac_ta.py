# -*- coding: utf-8 -*-
"""
sinh_dac_ta.py — Sinh toàn bộ file đặc tả của BTL nhóm 07 từ MỘT nguồn dữ liệu duy nhất.

Đề tài: Mô hình hóa và kiểm chứng hình thức quy trình thanh toán đơn hàng.

Đầu ra (cùng thư mục với script):
  thanhtoan_donhang_v0.xlsx / .json   — bản CÀI LỖI CÓ CHỦ ĐÍCH (18 trường hợp rủi ro)
  thanhtoan_donhang_v1.xlsx / .json   — bản ĐÃ SỬA
  mo_hinh_con/M1..M6_*.json           — 6 mô hình Kripke con (mỗi hành vi 1 file, nhiều luồng v0/v1)

File .json theo đúng định dạng dataset/specs/FORMAT.md của phần mềm (có thêm additional_flows
và một kiểm tra BMC vòng lặp `bmc_*` mà sheet Excel không biểu diễn được).
File .xlsx theo đúng 10 sheet của "Hướng dẫn tạo Spec file dạng excel.pdf".

Chạy:  python sinh_dac_ta.py        (cần `pip install openpyxl` để xuất .xlsx; JSON chỉ cần thư viện chuẩn)

Quy ước đặt tên:
  - Tên trạng thái: tiếng Việt không dấu, tiền tố theo hành vi (h1_.. h6_, m1_.. m6_).
  - Nhãn (atomic proposition): tiếng Anh không dấu, chữ thường — dùng trong công thức CTL.
  - Mọi công thức CTL ở dạng 1 toán tử + 1 nhãn (có thể phủ định), vì phần mềm CHỈ hỗ trợ dạng này.
"""
import json
import os

HERE = os.path.dirname(os.path.abspath(__file__))

# =====================================================================================
# 1. THÔNG TIN CHUNG
# =====================================================================================
MO_TA_CHUNG = (
    "Quy trình thanh toán trực tuyến 1 đơn hàng qua cổng thanh toán (có callback IPN), gồm 6 hành vi: "
    "H1 tiếp nhận yêu cầu (chép mã đơn vào bộ đệm), H2 tạo đơn (cấp phát mảng sản phẩm), "
    "H3 tính tổng tiền (số nguyên 32-bit, VND), H4 nhận callback từ cổng (chép mã giao dịch, xác thực chữ ký), "
    "H5 ghi log giao dịch, H6 đóng phiên (worker callback và worker timeout dùng chung phiên thanh toán)."
)
THONG_TIN = {
    "v0": ("thanhtoan_donhang_v0", MO_TA_CHUNG + " Bản v0 CÀI LỖI CÓ CHỦ ĐÍCH (18 trường hợp rủi ro bộ nhớ/tràn, "
           "trong đó 15 lỗi ẩn) để kiểm chứng khả năng phát hiện."),
    "v1": ("thanhtoan_donhang_v1", MO_TA_CHUNG + " Bản v1 ĐÃ SỬA: kiểm tra độ dài/miền giá trị và từ chối đầu vào xấu, "
           "kiểm tra NULL, calloc, tính tiền bằng int64 + hạn mức, xóa nhánh ẩn, giải phóng trên mọi nhánh, "
           "printf(\"%s\"), khóa phiên + trạng thái CLOSED, explicit_bzero trước khi free."),
}

# =====================================================================================
# 2. YÊU CẦU (chung cho v0 và v1 — yêu cầu không đổi, chỉ thiết kế thay đổi)
#    (ID, MoTa, DanhMuc, MucUuTien)
# =====================================================================================
YEU_CAU_CHUC_NANG = [
    ("REQ-01", "H1: Hàm phân tích mã đơn hàng (parse_order_id) không được ghi vượt biên bộ đệm, kể cả khi mã dài ĐÚNG 16 ký tự (C1.2, C1.3)", "memory", "HIGH"),
    ("REQ-02", "H1: Mã đơn hàng hợp lệ không được bị cắt cụt hoặc biến đổi khi lưu vào bộ đệm (C1.4)", "functional", "HIGH"),
    ("REQ-03", "H2: Kết quả cấp phát mảng sản phẩm (create_order_items) phải được kiểm tra NULL trước khi sử dụng (C2.2)", "memory", "HIGH"),
    ("REQ-04", "H2: Kích thước vùng nhớ cấp phát (số dòng x kích thước 1 dòng) không được tràn số hoặc âm (C2.3)", "safety", "HIGH"),
    ("REQ-05", "H2: Không được đọc trường dữ liệu của đơn hàng trước khi khởi tạo (C2.4)", "memory", "MEDIUM"),
    ("REQ-06", "H3: Tổng tiền đơn hàng không được tràn kiểu int32, kể cả khi cộng dồn nhiều dòng sản phẩm (C3.2, C3.3)", "safety", "HIGH"),
    ("REQ-07", "H3: Số lượng, đơn giá, giảm giá không được âm; tổng tiền phải thanh toán không được âm (C3.4)", "safety", "HIGH"),
    ("REQ-08", "H3: Hàm tính tiền trả góp không được gây lỗi không kiểm soát khi số kỳ hạn = 0 (C3.5)", "functional", "MEDIUM"),
    ("REQ-09", "H4: Hàm phân tích mã giao dịch từ callback (parse_gateway_txn_id) không được tràn bộ đệm (C4.2)", "memory", "HIGH"),
    ("REQ-10", "H4: Mọi callback phải được xác thực chữ ký; không được tồn tại nhánh ẩn bỏ qua kiểm tra (C4.3)", "safety", "HIGH"),
    ("REQ-11", "H5: Bản ghi log giao dịch (write_payment_log) phải được giải phóng trên mọi nhánh, kể cả nhánh lỗi (C5.2)", "memory", "MEDIUM"),
    ("REQ-12", "H5: Nội dung log có dữ liệu người dùng không được dùng làm chuỗi định dạng của printf (C5.3)", "safety", "MEDIUM"),
    ("REQ-13", "H6: Worker callback và worker timeout không được đồng thời giữ phiên thanh toán (mutual exclusion) (C6.2)", "safety", "HIGH"),
    ("REQ-14", "H6: Phiên thanh toán (close_payment_session) chỉ được giải phóng đúng 1 lần (C6.3)", "memory", "HIGH"),
    ("REQ-15", "H6: Không được truy cập phiên/bộ đệm thẻ (release_card_buffer) sau khi đã giải phóng (C6.4)", "memory", "HIGH"),
    ("REQ-16", "H6: Dữ liệu thẻ phải được xóa trắng bằng lệnh không bị trình biên dịch loại bỏ trước khi giải phóng (C6.5)", "safety", "HIGH"),
    ("REQ-17", "Toàn luồng: mọi đường xử lý thanh toán cuối cùng phải kết thúc (thành công hoặc từ chối an toàn), không dừng ở trạng thái không an toàn", "liveness", "HIGH"),
    ("REQ-18", "Toàn luồng: luồng thanh toán thành công phải đạt tới được (mô hình không tầm thường) và đầu vào xấu phải có nhánh từ chối", "liveness", "MEDIUM"),
]
YEU_CAU_PHI_CHUC_NANG = [
    ("NFR-01", "Hệ thống xử lý tối thiểu 200 yêu cầu thanh toán/giây trên phần cứng tham chiếu", "MEDIUM"),
    ("NFR-02", "Phản hồi callback cho cổng thanh toán trong tối đa 3 giây (quá hạn cổng sẽ gửi lại IPN, gây callback lặp)", "HIGH"),
    ("NFR-03", "Bộ nhớ tiến trình không tăng quá 5% sau 1.000.000 giao dịch liên tiếp (dấu hiệu rò rỉ bộ nhớ)", "HIGH"),
    ("NFR-04", "Phiên thanh toán tự hết hạn sau 15 phút nếu không nhận được callback (worker timeout)", "MEDIUM"),
    ("NFR-05", "Dữ liệu thẻ không được tồn tại trong bộ nhớ quá thời gian xử lý 1 giao dịch (tinh thần PCI-DSS)", "HIGH"),
    ("NFR-06", "Log giao dịch được lưu tối thiểu 12 tháng để đối soát", "LOW"),
]
RANG_BUOC = [
    ("CON-01", "Mã đơn hàng tối đa 16 ký tự ASCII (16 ký tự cần 17 byte kể cả '\\0')", "HIGH"),
    ("CON-02", "Mã giao dịch của cổng thanh toán tối đa 20 ký tự ASCII", "HIGH"),
    ("CON-03", "Mỗi đơn có 1–100 dòng sản phẩm; số lượng mỗi dòng 1–1000; đơn giá 0–50.000.000 VND; giá trị mỗi dòng tối đa 50.000.000 VND", "HIGH"),
    ("CON-04", "Số tiền tính bằng VND; bản v0 dùng int 32-bit có dấu (tối đa 2.147.483.647); hạn mức 1 đơn là 2.000.000.000 VND", "HIGH"),
    ("CON-05", "Kỳ hạn trả góp từ 1 đến 24 tháng", "MEDIUM"),
    ("CON-06", "Chỉ dùng thư viện chuẩn C; phiên thanh toán được bảo vệ bằng mutex nhẹ", "MEDIUM"),
    ("CON-07", "Bộ đệm số thẻ (PAN) 16 chữ số, chỉ tồn tại trong thời gian xử lý giao dịch", "HIGH"),
]

# =====================================================================================
# 3. MÔ HÌNH KRIPKE KẾT HỢP TOÀN LUỒNG (sheet MoHinhTrangThai / state_machine)
#    v0 và v1 có CÙNG tập trạng thái S và hàm nhãn L — chỉ khác quan hệ chuyển R.
#    Trạng thái rủi ro vẫn tồn tại ở v1 nhưng KHÔNG CÒN ĐƯỜNG ĐI tới (unreachable).
#    (ten, [nhan], ma_truong_hop, ghi_chu)
# =====================================================================================
TRANG_THAI = [
    ("start",              ["start", "safe"],                 "-",    "Bắt đầu xử lý 1 yêu cầu thanh toán"),
    # H1 — tiếp nhận yêu cầu, chép mã đơn vào order_id[16]
    ("h1_vao",             ["h1", "receiving"],               "H1",   "Nhận chuỗi mã đơn từ client"),
    ("h1_ok",              ["h1", "input_ok"],                "C1.1", "Mã đơn chép đủ, có '\\0' (an toàn)"),
    ("h1_tran_1_byte",     ["h1", "buffer_overflow", "off_by_one", "violation"], "C1.2", "Mã ĐÚNG 16 ký tự: strcpy ghi 17 byte, '\\0' đè sang biến kề (lỗi ẩn)"),
    ("h1_tran_bo_dem",     ["h1", "buffer_overflow", "violation"], "C1.3", "Mã > 16 ký tự: tràn bộ đệm stack"),
    # H2 — tạo đơn, cấp phát mảng sản phẩm
    ("h2_vao",             ["h2", "allocating"],              "H2",   "items = malloc(n * sizeof(Item))"),
    ("h2_ok",              ["h2", "allocated", "initialized"], "C2.1", "Cấp phát thành công và đã khởi tạo (an toàn)"),
    ("h2_null",            ["h2", "null_deref", "violation"], "C2.2", "malloc trả NULL nhưng vẫn ghi vào (lỗi ẩn: chỉ khi cạn bộ nhớ)"),
    ("h2_sai_kich_thuoc",  ["h2", "alloc_size_overflow", "heap_overflow", "violation"], "C2.3", "n*48 tràn int32 hoặc n âm: cấp phát sai kích thước rồi ghi tràn heap (lỗi ẩn)"),
    ("h2_doc_rac",         ["h2", "uninit_read", "violation"], "C2.4", "Đọc status khi chưa khởi tạo (lỗi ẩn)"),
    # H3 — tính tổng tiền (int32)
    ("h3_vao",             ["h3", "computing"],               "H3",   "Tính tổng = Σ đơn giá × SL + phí ship − giảm giá"),
    ("h3_ok",              ["h3", "amount_ok"],               "C3.1", "Tổng tiền nằm trong miền (an toàn)"),
    ("h3_tran_so",         ["h3", "int_overflow", "violation"], "C3.2/C3.3", "Tràn int32 khi nhân 1 dòng hoặc khi cộng dồn nhiều dòng (lỗi ẩn)"),
    ("h3_so_am",           ["h3", "negative_amount", "violation"], "C3.4", "SL/giảm giá âm hoặc giảm giá > tạm tính: tổng âm (lỗi ẩn)"),
    ("h3_chia_0",          ["h3", "div_by_zero", "violation"], "C3.5", "Trả góp với months = 0: chia cho 0"),
    # H4 — nhận callback IPN
    ("h4_vao",             ["h4", "callback_received"],       "H4",   "Nhận callback từ cổng thanh toán"),
    ("h4_ok",              ["h4", "callback_verified"],       "C4.1", "txn_id hợp lệ, chữ ký đúng (an toàn)"),
    ("h4_tu_choi_callback", ["h4", "callback_rejected"],      "C4.1", "Callback sai chữ ký bị bỏ qua (an toàn) — phiên sẽ do worker timeout đóng"),
    ("h4_tran_txn",        ["h4", "buffer_overflow", "violation"], "C4.2", "txn_id > 20 ký tự: tràn bộ đệm"),
    ("h4_nhanh_an",        ["h4", "hidden_bypass", "violation"], "C4.3", "Nhánh ẩn bỏ qua kiểm tra chữ ký (31337, 8080) (lỗi ẩn)"),
    # H5 — ghi log
    ("h5_vao",             ["h5", "logging"],                 "H5",   "e = malloc(LogEntry); ghi nội dung"),
    ("h5_ok",              ["h5", "log_freed"],               "C5.1", "Ghi log xong và free(e) (an toàn)"),
    ("h5_ro_ri",           ["h5", "mem_leak", "violation"],   "C5.2", "Nhánh lỗi return sớm, quên free (lỗi ẩn: rò rỉ tích lũy)"),
    ("h5_format_string",   ["h5", "format_string", "violation"], "C5.3", "printf(msg) với msg chứa %x/%n (lỗi ẩn)"),
    # H6 — đóng phiên, 2 worker dùng chung phiên (cụm 4 trạng thái khóa)
    ("h6_phien_ranh",      ["h6", "sess_free"],               "H6",   "Phiên đang mở, chưa worker nào giữ"),
    ("h6_cb_giu",          ["h6", "cb_locked"],               "C6.1", "Worker callback đang giữ phiên"),
    ("h6_to_giu",          ["h6", "to_locked"],               "C6.1", "Worker timeout (15 phút) đang giữ phiên"),
    ("h6_tranh_chap",      ["h6", "cb_locked", "to_locked", "race", "violation"], "C6.2", "Cả 2 worker cùng giữ phiên (race, lỗi ẩn)"),
    ("h6_double_free",     ["h6", "double_free", "violation"], "C6.3", "Cả 2 worker cùng free(session) (lỗi ẩn)"),
    ("h6_uaf",             ["h6", "use_after_free", "violation"], "C6.4", "Worker đến sau ghi vào phiên đã free (lỗi ẩn)"),
    ("h6_giai_phong",      ["h6", "releasing"],               "C6.1", "Giải phóng bộ đệm thẻ và phiên"),
    ("h6_con_so_the",      ["h6", "secret_residue", "violation"], "C6.5", "Số thẻ còn trong heap sau free (memset bị -O2 loại bỏ) (lỗi ẩn)"),
    # Kết thúc
    ("ket_thuc_ok",        ["flow_end", "payment_done", "mem_clean"], "-", "Kết thúc: đã xử lý thanh toán, bộ nhớ sạch"),
    ("ket_thuc_tu_choi",   ["flow_end", "input_rejected", "mem_clean"], "-", "Kết thúc: từ chối đầu vào xấu một cách an toàn"),
    ("dung_khong_an_toan", ["unsafe_halt"],                   "-",    "Trạng thái hút sau vi phạm (giống 'error location' trong BMC)"),
]

RUI_RO = [t[0] for t in TRANG_THAI if "violation" in t[1]]

def _chuyen_v0():
    R = [("start", "h1_vao"),
         ("h1_vao", "h1_ok"), ("h1_vao", "h1_tran_1_byte"), ("h1_vao", "h1_tran_bo_dem"),
         ("h1_ok", "h2_vao"),
         ("h2_vao", "h2_ok"), ("h2_vao", "h2_null"), ("h2_vao", "h2_sai_kich_thuoc"), ("h2_vao", "h2_doc_rac"),
         ("h2_ok", "h3_vao"),
         ("h3_vao", "h3_ok"), ("h3_vao", "h3_tran_so"), ("h3_vao", "h3_so_am"), ("h3_vao", "h3_chia_0"),
         ("h3_ok", "h4_vao"),
         ("h4_vao", "h4_ok"), ("h4_vao", "h4_tu_choi_callback"), ("h4_vao", "h4_tran_txn"), ("h4_vao", "h4_nhanh_an"),
         ("h4_ok", "h5_vao"), ("h4_tu_choi_callback", "h6_phien_ranh"),
         ("h5_vao", "h5_ok"), ("h5_vao", "h5_ro_ri"), ("h5_vao", "h5_format_string"),
         ("h5_ok", "h6_phien_ranh"),
         ("h6_phien_ranh", "h6_cb_giu"), ("h6_phien_ranh", "h6_to_giu"),
         ("h6_cb_giu", "h6_giai_phong"), ("h6_cb_giu", "h6_tranh_chap"),
         ("h6_to_giu", "h6_giai_phong"), ("h6_to_giu", "h6_tranh_chap"),
         ("h6_tranh_chap", "h6_double_free"), ("h6_tranh_chap", "h6_uaf"),
         ("h6_giai_phong", "ket_thuc_ok"), ("h6_giai_phong", "h6_con_so_the"),
         ("ket_thuc_ok", "ket_thuc_ok"), ("ket_thuc_tu_choi", "ket_thuc_tu_choi"),
         ("dung_khong_an_toan", "dung_khong_an_toan")]
    R += [(s, "dung_khong_an_toan") for s in RUI_RO if s != "h6_tranh_chap"]
    return R

def _chuyen_v1():
    R = [("start", "h1_vao"),
         ("h1_vao", "h1_ok"), ("h1_vao", "ket_thuc_tu_choi"),        # buf[17] + strnlen: len<=16 OK, >16 từ chối
         ("h1_ok", "h2_vao"),
         ("h2_vao", "h2_ok"), ("h2_vao", "ket_thuc_tu_choi"),        # n ngoài 1..100 hoặc NULL -> từ chối; calloc
         ("h2_ok", "h3_vao"),
         ("h3_vao", "h3_ok"), ("h3_vao", "ket_thuc_tu_choi"),        # int64 + hạn mức + kiểm tra miền + months>=1
         ("h3_ok", "h4_vao"),
         ("h4_vao", "h4_ok"), ("h4_vao", "h4_tu_choi_callback"),     # txn dài/sai chữ ký -> bỏ qua; xóa nhánh ẩn
         ("h4_ok", "h5_vao"), ("h4_tu_choi_callback", "h6_phien_ranh"),
         ("h5_vao", "h5_ok"),                                       # goto cleanup + printf("%s")
         ("h5_ok", "h6_phien_ranh"),
         ("h6_phien_ranh", "h6_cb_giu"), ("h6_phien_ranh", "h6_to_giu"),
         ("h6_cb_giu", "h6_giai_phong"), ("h6_to_giu", "h6_giai_phong"),   # mutex + status CLOSED
         ("h6_giai_phong", "ket_thuc_ok"),                          # explicit_bzero + free + NULL
         ("ket_thuc_ok", "ket_thuc_ok"), ("ket_thuc_tu_choi", "ket_thuc_tu_choi"),
         ("dung_khong_an_toan", "dung_khong_an_toan")]
    # Trạng thái rủi ro được GIỮ trong mô hình (không gian trạng thái không đổi) nhưng không còn cạnh đi tới.
    R += [(s, "dung_khong_an_toan") for s in RUI_RO if s != "h6_tranh_chap"]
    R += [("h6_tranh_chap", "h6_double_free"), ("h6_tranh_chap", "h6_uaf")]
    return R

CHUYEN = {"v0": _chuyen_v0(), "v1": _chuyen_v1()}

def state_machine(ver):
    trans = CHUYEN[ver]
    return {
        "states": [t[0] for t in TRANG_THAI],
        "initial": ["start"],
        "labeling": {t[0]: list(t[1]) for t in TRANG_THAI},
        "transitions": [list(x) for x in trans],
    }

# =====================================================================================
# 4. THUỘC TÍNH CTL (1 toán tử + 1 nhãn). (cong_thuc, ghi_chu)
# =====================================================================================
THUOC_TINH = [
    ("AG(!violation)",       "Safety tổng: không bao giờ đạt trạng thái vi phạm nào"),
    ("AG(!buffer_overflow)", "Safety C1.2, C1.3, C4.2 — CWE-120/787"),
    ("AG(!off_by_one)",      "Safety C1.2 — CWE-193 (lỗi ẩn: mã đúng 16 ký tự)"),
    ("AG(!null_deref)",      "Safety C2.2 — CWE-476/690"),
    ("AG(!heap_overflow)",   "Safety C2.3 — CWE-680/122"),
    ("AG(!uninit_read)",     "Safety C2.4 — CWE-457/908"),
    ("AG(!int_overflow)",    "Safety C3.2, C3.3 — CWE-190"),
    ("AG(!negative_amount)", "Safety C3.4 — CWE-191/839"),
    ("AG(!div_by_zero)",     "Safety C3.5 — CWE-369"),
    ("AG(!hidden_bypass)",   "Safety C4.3 — CWE-912/798"),
    ("AG(!mem_leak)",        "Safety C5.2 — CWE-401"),
    ("AG(!format_string)",   "Safety C5.3 — CWE-134"),
    ("AG(!race)",            "Safety C6.2 — CWE-362/667"),
    ("AG(!double_free)",     "Safety C6.3 — CWE-415"),
    ("AG(!use_after_free)",  "Safety C6.4 — CWE-416"),
    ("AG(!secret_residue)",  "Safety C6.5 — CWE-226/244/14"),
    ("AF(flow_end)",         "Liveness REQ-17: mọi đường đi cuối cùng đều kết thúc hợp lệ (LTL: F flow_end)"),
    ("EF(payment_done)",     "Khả đạt REQ-18: luồng thanh toán thành công tồn tại"),
    ("EF(input_rejected)",   "Khả đạt REQ-18: tồn tại nhánh từ chối đầu vào xấu (v0 KHÔNG có -> VI PHẠM)"),
    ("EG(!violation)",       "Tồn tại ít nhất 1 đường đi luôn an toàn"),
    ("EF(cb_locked)",        "Khả đạt của cụm khóa phiên (mô hình khóa không tầm thường)"),
]

LOAI_TRU = [  # dòng 1 là cặp được phần mềm kiểm tra tự động
    ("cb_locked", "to_locked"),
    ("violation", "flow_end"),
    ("payment_done", "input_rejected"),
    ("unsafe_halt", "mem_clean"),
    ("use_after_free", "mem_clean"),
]

# =====================================================================================
# 5. HÀM SINH MÃ (Ten, NgonNgu, Loai, ThamSo, KichThuocBuffer, Loi, GiaTriBiMat1, GiaTriBiMat2, GhiChu)
# =====================================================================================
HAM = {
    "v0": [
        ("parse_order_id",        "c", "buffer_copy",   "raw_order_id",  16, "buffer-overflow", None, None, "H1 — C1.2/C1.3 (CON-01)"),
        ("parse_gateway_txn_id",  "c", "buffer_copy",   "raw_txn_id",    20, "buffer-overflow", None, None, "H4 — C4.2 (CON-02)"),
        ("create_order_items",    "c", "alloc_release", "",             100, None,              None, None, "H2 — phần mềm báo SẠCH nhưng mô hình chỉ ra C2.2/C2.3/C2.4 (tầng 2)"),
        ("write_payment_log",     "c", "alloc_release", "",              32, "memory-leak",     None, None, "H5 — C5.2"),
        ("close_payment_session", "c", "alloc_release", "",               4, "double-free",     None, None, "H6 — C6.3"),
        ("release_card_buffer",   "c", "alloc_release", "",              16, "use-after-free",  None, None, "H6 — C6.4 (CON-07)"),
        ("add_shipping_fee",      "c", "arithmetic",    "subtotal, shipping_fee", None, None,   None, None, "H3 — trích AST cho Chương 3"),
        ("installment_amount",    "python", "divide",   "total, months", None, None,            None, None, "H3 — C3.5 (black-box fuzz)"),
        # Lưu ý: bộ fuzz của phần mềm đoán kiểu tham số theo TÊN — tên chứa code/key/name/data/msg/token...
        # bị coi là CHUỖI (làm black-box "crash" vì TypeError giả). Vì vậy đặt tên merchant_id, debug_pin (số nguyên).
        ("ipn_debug_bypass",      "python", "magic_value", "merchant_id, debug_pin", None, None, 31337, 8080, "H4 — C4.3 nhánh ẩn (concolic)"),
    ],
    "v1": [
        ("parse_order_id",        "c", "buffer_copy",   "raw_order_id",  17, None, None, None, "H1 — buffer 17 byte cho mã 16 ký tự (sửa C1.2)"),
        ("parse_gateway_txn_id",  "c", "buffer_copy",   "raw_txn_id",    21, None, None, None, "H4 — buffer 21 byte cho txn 20 ký tự"),
        ("create_order_items",    "c", "alloc_release", "",             100, None, None, None, "H2 — calloc + kiểm tra NULL + n trong 1..100"),
        ("write_payment_log",     "c", "alloc_release", "",              32, None, None, None, "H5 — goto cleanup: free trên mọi nhánh"),
        ("close_payment_session", "c", "alloc_release", "",               4, None, None, None, "H6 — free 1 lần, gán NULL, trạng thái CLOSED"),
        ("release_card_buffer",   "c", "alloc_release", "",              16, None, None, None, "H6 — explicit_bzero rồi free, không dùng lại"),
        ("add_shipping_fee",      "c", "arithmetic",    "subtotal, shipping_fee", None, None, None, None, "H3 — trích AST cho Chương 3"),
        # Hai hàm Python bị LOẠI ở v1: installment_amount được thay bằng kiểm tra months trong [1,24]
        # (chứng minh bằng KiemTraSoHoc/mô hình M3); ipn_debug_bypass là cửa sau -> bản sửa là XÓA.
    ],
}

# =====================================================================================
# 6. KIỂM TRA SỐ HỌC (Z3). (Ten, BienSoNguyen, DieuKien, BieuThucTran, GhiChu)
#    v0: DieuKien mô tả miền giá trị KHÔNG được kiểm tra -> kỳ vọng SAT (tồn tại phản ví dụ)
#    v1: DieuKien thêm ràng buộc kiểm tra đầu vào của bản sửa -> kỳ vọng UNSAT (chứng minh an toàn)
#    BieuThucTran luôn SAT ở cả 2 bản vì phần mềm luôn xét toàn miền int32 (đúng thiết kế, xem mục 3.9 hướng dẫn)
# =====================================================================================
SO_HOC = {
    "v0": [
        ("order_total_int32", "subtotal, shipping_fee",
         "subtotal >= 0 and subtotal <= 2147483647 and shipping_fee >= 0 and shipping_fee <= 500000 and subtotal + shipping_fee > 2147483647",
         "subtotal + shipping_fee", "C3.2 — cộng phí ship vào tạm tính gần trần int32"),
        ("line_amount_mul", "price, qty",
         "price >= 0 and price <= 50000000 and qty >= 1 and qty <= 1000 and price * qty > 2147483647",
         "price * qty", "C3.2 — phép nhân tràn TRƯỚC khi kịp kiểm tra hạn mức dòng"),
        ("accumulated_total", "total_before, line_total",
         "total_before >= 0 and total_before <= 2147483647 and line_total >= 0 and line_total <= 50000000 and total_before + line_total > 2147483647",
         "total_before + line_total", "C3.3 — bước cộng dồn: mỗi dòng hợp lệ nhưng tổng vượt trần"),
        ("alloc_bytes", "n_items, item_size",
         "n_items >= 1 and item_size == 48 and n_items * item_size > 2147483647",
         "n_items * item_size", "C2.3 — kích thước cấp phát tràn khi n không bị giới hạn"),
        ("negative_total", "subtotal, discount",
         # Lưu ý: bộ phân tích biểu thức của phần mềm KHÔNG nhận số âm dạng "-2000000000"
         # -> viết tương đương "discount + 2000000000 >= 0" (discount >= -2.000.000.000)
         "subtotal >= 0 and subtotal <= 2000000000 and discount + 2000000000 >= 0 and discount <= 2000000000 and subtotal - discount < 0",
         "subtotal - discount", "C3.4 — giảm giá âm hoặc lớn hơn tạm tính"),
    ],
    "v1": [
        ("order_total_int32", "subtotal, shipping_fee",
         "subtotal >= 0 and subtotal <= 2000000000 and shipping_fee >= 0 and shipping_fee <= 500000 and subtotal + shipping_fee > 2147483647",
         "subtotal + shipping_fee", "v1: hạn mức đơn 2 tỷ -> UNSAT"),
        ("line_amount_mul", "price, qty",
         "price >= 0 and price <= 50000000 and qty >= 1 and qty <= 1000 and price * qty <= 50000000 and price * qty > 2147483647",
         "price * qty", "v1: nhân bằng int64 rồi kiểm tra hạn mức dòng trước khi lưu -> UNSAT"),
        ("accumulated_total", "total_before, line_total",
         "total_before >= 0 and line_total >= 0 and line_total <= 50000000 and total_before <= 2000000000 - line_total and total_before + line_total > 2147483647",
         "total_before + line_total", "v1: kiểm tra hạn mức TRƯỚC mỗi lần cộng (bước quy nạp) -> UNSAT"),
        ("alloc_bytes", "n_items, item_size",
         "n_items >= 1 and n_items <= 100 and item_size == 48 and n_items * item_size > 2147483647",
         "n_items * item_size", "v1: n trong 1..100 -> UNSAT"),
        ("negative_total", "subtotal, discount",
         "subtotal >= 0 and subtotal <= 2000000000 and discount >= 0 and discount <= subtotal and subtotal - discount < 0",
         "subtotal - discount", "v1: 0 <= giảm giá <= tạm tính -> UNSAT"),
    ],
}
# BMC vòng lặp (chỉ có trong file JSON — Excel không có cột bmc_*): cộng dồn 45.000.000 VND/dòng
BMC_VONG_LAP = {
    "v0": {"name": "bmc_cong_don_don_hang", "int_vars": ["i", "total"],
           "bmc_init": {"i": 0, "total": 0}, "bmc_step": "i = i + 1; total = total + 45000000",
           "bmc_cond": "i < 100", "bmc_assert": "total <= 2147483647", "bmc_k": 100},
    "v1": {"name": "bmc_cong_don_don_hang", "int_vars": ["i", "total"],
           "bmc_init": {"i": 0, "total": 0}, "bmc_step": "i = i + 1; total = total + 45000000",
           "bmc_cond": "i < 100 and total + 45000000 <= 2000000000", "bmc_assert": "total <= 2147483647", "bmc_k": 100},
}

# =====================================================================================
# 7. TEST CASE truy vết (ID, YeuCauLienQuan, TieuDe, DieuKienTienQuyet, [CacBuoc], KetQuaMongDoi, MucUuTien)
#    Test case được viết TỪ kết quả kiểm chứng (counterexample), không viết trước.
# =====================================================================================
TEST_CASE = [
    ("TC-01", "REQ-01, CON-01", "Mã đơn dài đúng 16 ký tự và dài 100 ký tự",
     "parse_order_id sinh từ đặc tả (v0: Loi=buffer-overflow, buffer 16)",
     ["Sinh mã C cho parse_order_id", "Chạy check_c_source() (Chương 2)", "Kiểm tra AG(!buffer_overflow), AG(!off_by_one) trên mô hình"],
     "v0: dangerous-function strcpy (CWE-120), 2 công thức VI PHẠM; v1: 0 phát hiện, công thức ĐÚNG", "HIGH"),
    ("TC-02", "REQ-02, CON-01", "Bản sửa của phần mềm có cắt cụt mã đơn 16 ký tự không",
     "Mô hình con M1, luồng 'M1-v1a' (strncpy 15 + '\\0')",
     ["Nạp mo_hinh_con/M1_tiep_nhan_yeu_cau.json vào tab Ch1", "Chọn luồng M1-v1a, kiểm tra AG(!truncated)", "So sánh với luồng M1-v1b"],
     "M1-v1a: AG(!truncated) VI PHẠM (mã bị cắt, phần mềm không báo); M1-v1b: ĐÚNG", "HIGH"),
    ("TC-03", "REQ-03, REQ-04, REQ-05", "Cấp phát mảng sản phẩm",
     "create_order_items (alloc_release, không cài lỗi); mô hình con M2",
     ["Chạy check_c_source() trên mã sinh", "Kiểm tra M2-v0: AG(!null_deref), AG(!heap_overflow), AG(!uninit_read)", "Chạy KiemTraSoHoc alloc_bytes"],
     "Phần mềm quét mã: 0 phát hiện; mô hình M2-v0: 3 công thức VI PHẠM; alloc_bytes SAT (v0) / UNSAT (v1)", "HIGH"),
    ("TC-04", "REQ-06, CON-03, CON-04", "Tổng tiền tràn int32",
     "KiemTraSoHoc order_total_int32, line_amount_mul, accumulated_total; BMC vòng lặp trong file JSON",
     ["Chạy check_satisfiability() trên DieuKien", "Chạy BMC bmc_cong_don_don_hang (k = 100)"],
     "v0: SAT kèm phản ví dụ, BMC vi phạm tại bước 48 (total = 2.160.000.000); v1: UNSAT, BMC không vi phạm", "HIGH"),
    ("TC-05", "REQ-07, CON-03", "Số tiền âm do giảm giá",
     "KiemTraSoHoc negative_total",
     ["Chạy check_satisfiability() trên DieuKien"],
     "v0: SAT (vd discount > subtotal); v1: UNSAT", "HIGH"),
    ("TC-06", "REQ-08, CON-05", "Trả góp với số kỳ hạn = 0",
     "Hàm Python installment_amount(total, months) (kind = divide) — chỉ có ở v0",
     ["Chạy blackbox_fuzz() 300 lần"],
     "v0: tìm được ít nhất 1 crash ZeroDivisionError (CWE-248/369); v1: hàm đã thay bằng kiểm tra 1 <= months <= 24", "MEDIUM"),
    ("TC-07", "REQ-09, CON-02", "Mã giao dịch callback dài hơn 20 ký tự",
     "parse_gateway_txn_id (v0: buffer 20, Loi=buffer-overflow)",
     ["Sinh mã C", "Chạy check_c_source()"],
     "v0: dangerous-function (CWE-120); v1: 0 phát hiện", "HIGH"),
    ("TC-08", "REQ-10", "Nhánh ẩn bỏ qua xác thực chữ ký",
     "Hàm Python ipn_debug_bypass(merchant_id, debug_pin), magic_values = (31337, 8080)",
     ["Chạy black-box fuzz trước (kỳ vọng không tìm ra)", "Chạy concolic_fuzz()"],
     "Black-box: 0 lần chạm nhánh; concolic: tìm đúng (31337, 8080) sau vài lần thử (CWE-912)", "HIGH"),
    ("TC-09", "REQ-11, NFR-03", "Rò rỉ bộ nhớ khi ghi log lỗi",
     "write_payment_log (v0: Loi=memory-leak)",
     ["Sinh mã C", "Chạy check_c_source()"],
     "v0: memory-leak (CWE-401); v1: 0 phát hiện", "MEDIUM"),
    ("TC-10", "REQ-12", "Chuỗi định dạng trong log",
     "Mô hình con M5 (mo_hinh_con/M5_ghi_log.json) — khuôn mẫu sinh mã của phần mềm không tạo được lỗi này",
     ["Nạp M5_ghi_log.json vào tab Ch1", "Chọn luồng M5-v0, kiểm tra AG(!format_string)", "Chọn luồng M5-v1, kiểm tra lại"],
     "M5-v0: AG(!format_string) VI PHẠM (CWE-134); M5-v1: ĐÚNG", "MEDIUM"),
    ("TC-11", "REQ-13, NFR-04", "Callback và timeout cùng đóng phiên",
     "Cụm trạng thái h6_* trong mô hình kết hợp; LoaiTruLanNhau dòng 1 (cb_locked, to_locked)",
     ["Worker timeout kích hoạt đúng lúc worker callback đang giữ phiên", "Chạy check_mutual_exclusion(model, 'cb_locked', 'to_locked')"],
     "v0: Mutual exclusion VI PHẠM tại h6_tranh_chap (CWE-667); v1: ĐÚNG", "HIGH"),
    ("TC-12", "REQ-14, REQ-15", "Double free và use-after-free khi đóng phiên",
     "close_payment_session (double-free), release_card_buffer (use-after-free)",
     ["Gộp mã C của 2 hàm", "Chạy check_c_source()"],
     "v0: double-free (CWE-415) và use-after-free (CWE-416); v1: 0 phát hiện", "HIGH"),
    ("TC-13", "REQ-16, NFR-05, CON-07", "Số thẻ còn trong bộ nhớ sau khi giải phóng",
     "Mô hình con M6 (mo_hinh_con/M6_dong_phien.json)",
     ["Nạp M6_dong_phien.json vào tab Ch1", "Chọn luồng M6-v0, kiểm tra AG(!secret_residue)", "Chọn luồng M6-v1, kiểm tra lại"],
     "M6-v0: VI PHẠM (CWE-226/244, CWE-14); M6-v1 (explicit_bzero trước free): ĐÚNG", "HIGH"),
    ("TC-14", "REQ-17, REQ-18", "Toàn luồng thanh toán",
     "Mô hình kết hợp (sheet MoHinhTrangThai)",
     ["Chạy '1) Sinh Kripke + kiểm chứng (Ch1)'", "Kiểm tra deadlock, AF(flow_end), EF(payment_done), EF(input_rejected)"],
     "Không deadlock; v0: AF(flow_end) và EF(input_rejected) VI PHẠM; v1: tất cả ĐÚNG", "HIGH"),
]

# =====================================================================================
# 8. SÁU MÔ HÌNH KRIPKE CON (tab Ch1 SecLabFramework — mỗi luồng được kiểm chứng riêng)
#    Mức chi tiết: mỗi trạng thái ~ 1 câu lệnh / 1 điều kiện của mã C tương ứng.
# =====================================================================================
def _flow(desc, states, labels, trans, init):
    used = list(dict.fromkeys(states))
    return {"description": desc, "states": used, "initial": [init],
            "labeling": {s: labels[s] for s in used}, "transitions": [list(t) for t in trans]}

def _self(*ss):
    return [(s, s) for s in ss]

MO_HINH_CON = {}

# ---- M1: H1 tiếp nhận yêu cầu — chép mã đơn vào bộ đệm ----
L1 = {"m1_nhan_chuoi": ["receiving"], "m1_len_nho": ["len_lt_cap"], "m1_len_bang_16": ["len_eq_cap"],
      "m1_len_lon": ["len_gt_cap"], "m1_chep_dung": ["copied_ok", "done"],
      "m1_tran_1_byte": ["buffer_overflow", "off_by_one", "violation"], "m1_tran_nhieu_byte": ["buffer_overflow", "violation"],
      "m1_cat_cut": ["truncated", "violation"], "m1_tu_choi": ["input_rejected", "done"]}
_vao1 = [("m1_nhan_chuoi", "m1_len_nho"), ("m1_nhan_chuoi", "m1_len_bang_16"), ("m1_nhan_chuoi", "m1_len_lon")]
MO_HINH_CON["M1_tiep_nhan_yeu_cau"] = {
    "title": "M1 — H1 Tiếp nhận yêu cầu: chép mã đơn (tối đa 16 ký tự) vào bộ đệm",
    "description": "Phân lớp độ dài L của mã đơn so với dung lượng bộ đệm N=16: L<16, L=16, L>16 phủ MỌI độ dài. "
                   "So sánh 3 phiên bản: v0 strcpy, v1a bản 'an toàn' do phần mềm sinh (strncpy 15 + '\\0'), v1b bản sửa của nhóm.",
    "flows": {
        "M1-v0: strcpy vào order_id[16] (CÓ LỖI)": _flow(
            "char order_id[16]; strcpy(order_id, src);",
            ["m1_nhan_chuoi", "m1_len_nho", "m1_len_bang_16", "m1_len_lon", "m1_chep_dung", "m1_tran_1_byte", "m1_tran_nhieu_byte"], L1,
            _vao1 + [("m1_len_nho", "m1_chep_dung"), ("m1_len_bang_16", "m1_tran_1_byte"), ("m1_len_lon", "m1_tran_nhieu_byte")]
            + _self("m1_chep_dung", "m1_tran_1_byte", "m1_tran_nhieu_byte"), "m1_nhan_chuoi"),
        "M1-v1a: bản 'an toàn' phần mềm sinh — strncpy(buf,src,15) (CẮT CỤT)": _flow(
            "char buffer[16]; strncpy(buffer, src, sizeof(buffer)-1); buffer[15] = '\\0';",
            ["m1_nhan_chuoi", "m1_len_nho", "m1_len_bang_16", "m1_len_lon", "m1_chep_dung", "m1_cat_cut"], L1,
            _vao1 + [("m1_len_nho", "m1_chep_dung"), ("m1_len_bang_16", "m1_cat_cut"), ("m1_len_lon", "m1_cat_cut")]
            + _self("m1_chep_dung", "m1_cat_cut"), "m1_nhan_chuoi"),
        "M1-v1b: bản sửa của nhóm — buf[17] + strnlen + từ chối": _flow(
            "char order_id[17]; if (strnlen(src, 17) > 16) return REJECT; memcpy(order_id, src, len + 1);",
            ["m1_nhan_chuoi", "m1_len_nho", "m1_len_bang_16", "m1_len_lon", "m1_chep_dung", "m1_tu_choi"], L1,
            _vao1 + [("m1_len_nho", "m1_chep_dung"), ("m1_len_bang_16", "m1_chep_dung"), ("m1_len_lon", "m1_tu_choi")]
            + _self("m1_chep_dung", "m1_tu_choi"), "m1_nhan_chuoi"),
    },
    "suggested_ctl_formulas": ["AG(!buffer_overflow)", "AG(!off_by_one)", "AG(!truncated)", "AF(done)", "EF(len_eq_cap)"],
}

# ---- M2: H2 tạo đơn — vòng đời con trỏ mảng sản phẩm ----
L2 = {"m2_nhan_n": ["receiving"], "m2_n_hop_le": ["size_ok"], "m2_n_bat_thuong": ["size_wrong"],
      "m2_malloc_null": ["ptr_null"], "m2_cap_phat_chua_kt": ["allocated", "uninitialized"],
      "m2_da_khoi_tao": ["allocated", "initialized"], "m2_null_deref": ["null_deref", "violation"],
      "m2_tran_heap": ["heap_overflow", "violation"], "m2_doc_rac": ["uninit_read", "violation"],
      "m2_su_dung": ["in_use", "done"], "m2_tu_choi": ["input_rejected", "done"]}
MO_HINH_CON["M2_tao_don_cap_phat"] = {
    "title": "M2 — H2 Tạo đơn: items = malloc(n * sizeof(Item)), sizeof(Item) = 48",
    "description": "Vòng đời con trỏ: NULL / đã cấp phát chưa khởi tạo / đã khởi tạo; cộng thêm phân lớp n (hợp lệ 1..100 hoặc bất thường).",
    "flows": {
        "M2-v0: malloc không kiểm tra (CÓ LỖI)": _flow(
            "int bytes = n * (int)sizeof(Item); Item *items = malloc(bytes); items[0].qty = q; if (items[i].status == PAID) ...",
            ["m2_nhan_n", "m2_n_hop_le", "m2_n_bat_thuong", "m2_malloc_null", "m2_cap_phat_chua_kt", "m2_da_khoi_tao",
             "m2_null_deref", "m2_tran_heap", "m2_doc_rac", "m2_su_dung"], L2,
            [("m2_nhan_n", "m2_n_hop_le"), ("m2_nhan_n", "m2_n_bat_thuong"),
             ("m2_n_hop_le", "m2_malloc_null"), ("m2_n_hop_le", "m2_cap_phat_chua_kt"),
             ("m2_n_bat_thuong", "m2_tran_heap"), ("m2_malloc_null", "m2_null_deref"),
             ("m2_cap_phat_chua_kt", "m2_doc_rac"), ("m2_cap_phat_chua_kt", "m2_da_khoi_tao"),
             ("m2_da_khoi_tao", "m2_su_dung")] + _self("m2_null_deref", "m2_tran_heap", "m2_doc_rac", "m2_su_dung"),
            "m2_nhan_n"),
        "M2-v1: kiểm tra n, calloc, kiểm tra NULL": _flow(
            "if (n < 1 || n > 100) return REJECT; Item *items = calloc(n, sizeof(Item)); if (!items) return REJECT;",
            ["m2_nhan_n", "m2_n_hop_le", "m2_malloc_null", "m2_da_khoi_tao", "m2_su_dung", "m2_tu_choi"], L2,
            [("m2_nhan_n", "m2_n_hop_le"), ("m2_nhan_n", "m2_tu_choi"),
             ("m2_n_hop_le", "m2_malloc_null"), ("m2_n_hop_le", "m2_da_khoi_tao"),
             ("m2_malloc_null", "m2_tu_choi"), ("m2_da_khoi_tao", "m2_su_dung")] + _self("m2_su_dung", "m2_tu_choi"),
            "m2_nhan_n"),
    },
    "suggested_ctl_formulas": ["AG(!null_deref)", "AG(!heap_overflow)", "AG(!uninit_read)", "AF(done)", "EF(ptr_null)"],
}

# ---- M3: H3 tính tổng tiền (int32) ----
L3 = {"m3_nhan_du_lieu": ["receiving"], "m3_dau_vao_hop_le": ["input_ok"], "m3_dau_vao_am": ["negative_input"],
      "m3_ky_han_0": ["zero_divisor"], "m3_nhan_dong": ["multiplying"], "m3_cong_don": ["accumulating"],
      "m3_tong_hop_le": ["amount_ok", "done"], "m3_tran_nhan": ["int_overflow", "violation"],
      "m3_tran_cong_don": ["int_overflow", "violation"], "m3_tong_am": ["negative_amount", "violation"],
      "m3_chia_0": ["div_by_zero", "violation"], "m3_tu_choi": ["input_rejected", "done"]}
MO_HINH_CON["M3_tinh_tong_tien"] = {
    "title": "M3 — H3 Tính tổng tiền: total = Σ price*qty + ship − discount (int32, VND)",
    "description": "Phân lớp miền số nguyên: trong miền / tràn trên khi nhân / tràn khi cộng dồn / âm / mẫu số bằng 0.",
    "flows": {
        "M3-v0: int32, không kiểm tra miền (CÓ LỖI)": _flow(
            "int line = price * qty; total += line; total -= discount; per_month = total / months;",
            ["m3_nhan_du_lieu", "m3_dau_vao_hop_le", "m3_dau_vao_am", "m3_ky_han_0", "m3_nhan_dong", "m3_cong_don",
             "m3_tong_hop_le", "m3_tran_nhan", "m3_tran_cong_don", "m3_tong_am", "m3_chia_0"], L3,
            [("m3_nhan_du_lieu", "m3_dau_vao_hop_le"), ("m3_nhan_du_lieu", "m3_dau_vao_am"), ("m3_nhan_du_lieu", "m3_ky_han_0"),
             ("m3_dau_vao_hop_le", "m3_nhan_dong"), ("m3_nhan_dong", "m3_cong_don"), ("m3_nhan_dong", "m3_tran_nhan"),
             ("m3_cong_don", "m3_tong_hop_le"), ("m3_cong_don", "m3_tran_cong_don"),
             ("m3_dau_vao_am", "m3_tong_am"), ("m3_ky_han_0", "m3_chia_0")]
            + _self("m3_tong_hop_le", "m3_tran_nhan", "m3_tran_cong_don", "m3_tong_am", "m3_chia_0"), "m3_nhan_du_lieu"),
        "M3-v1: kiểm tra miền + int64 + hạn mức": _flow(
            "if (qty<1||qty>1000||price<0||price>50000000||discount<0||months<1||months>24) REJECT; "
            "int64_t line = (int64_t)price*qty; if (line > 50000000) REJECT; if (total + line > 2000000000) REJECT;",
            ["m3_nhan_du_lieu", "m3_dau_vao_hop_le", "m3_nhan_dong", "m3_cong_don", "m3_tong_hop_le", "m3_tu_choi"], L3,
            [("m3_nhan_du_lieu", "m3_dau_vao_hop_le"), ("m3_nhan_du_lieu", "m3_tu_choi"),
             ("m3_dau_vao_hop_le", "m3_nhan_dong"), ("m3_nhan_dong", "m3_cong_don"), ("m3_nhan_dong", "m3_tu_choi"),
             ("m3_cong_don", "m3_tong_hop_le"), ("m3_cong_don", "m3_tu_choi")] + _self("m3_tong_hop_le", "m3_tu_choi"),
            "m3_nhan_du_lieu"),
    },
    "suggested_ctl_formulas": ["AG(!int_overflow)", "AG(!negative_amount)", "AG(!div_by_zero)", "AF(done)", "EF(amount_ok)"],
}

# ---- M4: H4 nhận callback IPN ----
L4 = {"m4_nhan_callback": ["receiving"], "m4_txn_hop_le": ["len_ok"], "m4_txn_qua_dai": ["len_gt_cap"],
      "m4_tran_txn": ["buffer_overflow", "violation"], "m4_kiem_chu_ky": ["verifying"],
      "m4_nhanh_an": ["hidden_bypass", "violation"], "m4_chu_ky_dung": ["verified"],
      "m4_tu_choi": ["callback_rejected", "done"], "m4_chap_nhan": ["accepted", "done"]}
MO_HINH_CON["M4_nhan_callback"] = {
    "title": "M4 — H4 Nhận callback (IPN) từ cổng thanh toán",
    "description": "Dữ liệu đến từ BÊN NGOÀI (attack surface): độ dài txn_id, chữ ký, và nhánh ẩn do lập trình viên bỏ quên.",
    "flows": {
        "M4-v0: strcpy txn_id[20] + nhánh debug (CÓ LỖI)": _flow(
            "strcpy(txn, cb->txn_id); if (merchant == 31337 && key == 8080) goto accept; if (!verify_hmac(cb)) reject;",
            ["m4_nhan_callback", "m4_txn_hop_le", "m4_txn_qua_dai", "m4_tran_txn", "m4_kiem_chu_ky", "m4_nhanh_an",
             "m4_chu_ky_dung", "m4_tu_choi", "m4_chap_nhan"], L4,
            [("m4_nhan_callback", "m4_txn_hop_le"), ("m4_nhan_callback", "m4_txn_qua_dai"), ("m4_txn_qua_dai", "m4_tran_txn"),
             ("m4_txn_hop_le", "m4_kiem_chu_ky"), ("m4_txn_hop_le", "m4_nhanh_an"),
             ("m4_kiem_chu_ky", "m4_chu_ky_dung"), ("m4_kiem_chu_ky", "m4_tu_choi"), ("m4_chu_ky_dung", "m4_chap_nhan")]
            + _self("m4_tran_txn", "m4_nhanh_an", "m4_tu_choi", "m4_chap_nhan"), "m4_nhan_callback"),
        "M4-v1: kiểm tra độ dài, xóa nhánh ẩn": _flow(
            "if (strnlen(cb->txn_id, 21) > 20) reject; if (!verify_hmac(cb)) reject; accept;",
            ["m4_nhan_callback", "m4_txn_hop_le", "m4_kiem_chu_ky", "m4_chu_ky_dung", "m4_tu_choi", "m4_chap_nhan"], L4,
            [("m4_nhan_callback", "m4_txn_hop_le"), ("m4_nhan_callback", "m4_tu_choi"), ("m4_txn_hop_le", "m4_kiem_chu_ky"),
             ("m4_kiem_chu_ky", "m4_chu_ky_dung"), ("m4_kiem_chu_ky", "m4_tu_choi"), ("m4_chu_ky_dung", "m4_chap_nhan")]
            + _self("m4_tu_choi", "m4_chap_nhan"), "m4_nhan_callback"),
    },
    "suggested_ctl_formulas": ["AG(!buffer_overflow)", "AG(!hidden_bypass)", "AF(done)", "EF(accepted)", "EF(callback_rejected)"],
}

# ---- M5: H5 ghi log ----
L5 = {"m5_bat_dau": ["receiving"], "m5_cap_phat": ["allocated"], "m5_ghi_noi_dung": ["allocated", "in_use"],
      "m5_in_an_toan": ["allocated", "printed"], "m5_format_string": ["format_string", "violation"],
      "m5_loi_ghi_dia": ["allocated", "error_path"], "m5_ro_ri": ["mem_leak", "violation"],
      "m5_giai_phong": ["freed", "done"]}
MO_HINH_CON["M5_ghi_log"] = {
    "title": "M5 — H5 Ghi log giao dịch: e = malloc(sizeof(LogEntry))",
    "description": "Vòng đời con trỏ bản ghi log trên nhánh thường và nhánh lỗi; cách in nội dung log.",
    "flows": {
        "M5-v0: printf(msg) + return sớm (CÓ LỖI)": _flow(
            "e = malloc(...); snprintf(e->msg, ...); printf(e->msg); if (fwrite(...) != 1) return -1; free(e);",
            ["m5_bat_dau", "m5_cap_phat", "m5_ghi_noi_dung", "m5_in_an_toan", "m5_format_string", "m5_loi_ghi_dia", "m5_ro_ri", "m5_giai_phong"], L5,
            [("m5_bat_dau", "m5_cap_phat"), ("m5_cap_phat", "m5_ghi_noi_dung"),
             ("m5_ghi_noi_dung", "m5_in_an_toan"), ("m5_ghi_noi_dung", "m5_format_string"), ("m5_ghi_noi_dung", "m5_loi_ghi_dia"),
             ("m5_in_an_toan", "m5_giai_phong"), ("m5_loi_ghi_dia", "m5_ro_ri")]
            + _self("m5_format_string", "m5_ro_ri", "m5_giai_phong"), "m5_bat_dau"),
        "M5-v1: printf(\"%s\") + goto cleanup": _flow(
            "printf(\"%s\", e->msg); if (fwrite(...) != 1) { rc = -1; goto cleanup; } cleanup: free(e); e = NULL;",
            ["m5_bat_dau", "m5_cap_phat", "m5_ghi_noi_dung", "m5_in_an_toan", "m5_loi_ghi_dia", "m5_giai_phong"], L5,
            [("m5_bat_dau", "m5_cap_phat"), ("m5_cap_phat", "m5_ghi_noi_dung"),
             ("m5_ghi_noi_dung", "m5_in_an_toan"), ("m5_ghi_noi_dung", "m5_loi_ghi_dia"),
             ("m5_in_an_toan", "m5_giai_phong"), ("m5_loi_ghi_dia", "m5_giai_phong")] + _self("m5_giai_phong"), "m5_bat_dau"),
    },
    "suggested_ctl_formulas": ["AG(!mem_leak)", "AG(!format_string)", "AF(freed)", "EF(error_path)"],
}

# ---- M6: H6 đóng phiên — 2 worker dùng chung phiên + bộ đệm số thẻ ----
L6 = {"m6_phien_mo": ["sess_free"], "m6_cb_giu": ["cb_locked"], "m6_to_giu": ["to_locked"],
      "m6_tranh_chap": ["cb_locked", "to_locked", "race", "violation"],
      "m6_double_free": ["double_free", "violation"], "m6_uaf": ["use_after_free", "violation"],
      "m6_giai_phong": ["releasing"], "m6_con_so_the": ["secret_residue", "violation"],
      "m6_xoa_trang": ["wiped"], "m6_da_dong": ["session_freed", "done"],
      "m6_den_sau_bo_qua": ["late_ignored", "session_freed", "done"]}
MO_HINH_CON["M6_dong_phien"] = {
    "title": "M6 — H6 Đóng phiên: worker callback và worker timeout dùng chung PaymentSession *s",
    "description": "Xét MỌI thứ tự xen kẽ của 2 worker (interleaving) và vòng đời của s, card: giữ khóa / giải phóng / dùng lại.",
    "flows": {
        "M6-v0: không khóa, memset trước free (CÓ LỖI)": _flow(
            "on_callback(){ s->status = PAID; close(s); }  on_timeout(){ s->status = CANCELLED; close(s); }  "
            "close(s){ memset(s->card,0,16); free(s->card); free(s); }",
            ["m6_phien_mo", "m6_cb_giu", "m6_to_giu", "m6_tranh_chap", "m6_double_free", "m6_uaf",
             "m6_giai_phong", "m6_con_so_the", "m6_da_dong"], L6,
            [("m6_phien_mo", "m6_cb_giu"), ("m6_phien_mo", "m6_to_giu"),
             ("m6_cb_giu", "m6_giai_phong"), ("m6_cb_giu", "m6_tranh_chap"),
             ("m6_to_giu", "m6_giai_phong"), ("m6_to_giu", "m6_tranh_chap"),
             ("m6_tranh_chap", "m6_double_free"), ("m6_tranh_chap", "m6_uaf"),
             ("m6_giai_phong", "m6_con_so_the"), ("m6_giai_phong", "m6_da_dong"),
             ("m6_da_dong", "m6_da_dong"), ("m6_da_dong", "m6_uaf")]          # callback đến muộn ghi vào phiên đã free
            + _self("m6_double_free", "m6_uaf", "m6_con_so_the"), "m6_phien_mo"),
        "M6-v1: mutex + status CLOSED + explicit_bzero": _flow(
            "lock(s->mtx); if (s->status == CLOSED) { unlock; return; } ... explicit_bzero(s->card,16); free(s->card); "
            "s->status = CLOSED; unlock; (worker đến sau thấy CLOSED thì bỏ qua)",
            ["m6_phien_mo", "m6_cb_giu", "m6_to_giu", "m6_giai_phong", "m6_xoa_trang", "m6_da_dong", "m6_den_sau_bo_qua"], L6,
            [("m6_phien_mo", "m6_cb_giu"), ("m6_phien_mo", "m6_to_giu"),
             ("m6_cb_giu", "m6_giai_phong"), ("m6_to_giu", "m6_giai_phong"),
             ("m6_giai_phong", "m6_xoa_trang"), ("m6_xoa_trang", "m6_da_dong"),
             ("m6_da_dong", "m6_den_sau_bo_qua")] + _self("m6_den_sau_bo_qua"), "m6_phien_mo"),
    },
    "suggested_ctl_formulas": ["AG(!race)", "AG(!double_free)", "AG(!use_after_free)", "AG(!secret_residue)", "AF(session_freed)", "EF(cb_locked)"],
    "suggested_mutex_atoms": ["cb_locked", "to_locked"],
}


# =====================================================================================
# 9. DỰNG ĐẶC TẢ JSON (định dạng FORMAT.md) VÀ GHI FILE
# =====================================================================================
def _ham_json(h):
    ten, ngon_ngu, loai, tham_so, buf, loi, m1, m2, _ = h
    d = {"name": ten, "language": ngon_ngu, "kind": loai,
         "params": [p.strip() for p in tham_so.split(",") if p.strip()]}
    if loai in ("buffer_copy", "alloc_release"):
        d["buffer_size"] = buf
        d["bug"] = loi
    if loai == "magic_value":
        d["magic_values"] = [m1, m2]
    return d

def dac_ta_json(ver):
    ten, mo_ta = THONG_TIN[ver]
    reqs = [{"id": i, "text": t, "category": c, "priority": p} for i, t, c, p in YEU_CAU_CHUC_NANG]
    reqs += [{"id": i, "text": t, "category": "non_functional", "priority": p} for i, t, p in YEU_CAU_PHI_CHUC_NANG]
    reqs += [{"id": i, "text": t, "category": "constraint", "priority": p} for i, t, p in RANG_BUOC]
    arith = [{"name": n, "int_vars": [v.strip() for v in bien.split(",")], "condition": dk, "overflow_expr": bt}
             for n, bien, dk, bt, _ in SO_HOC[ver]]
    arith.append(dict(BMC_VONG_LAP[ver]))
    flows = {}
    for key, mh in MO_HINH_CON.items():
        for ten_luong, luong in mh["flows"].items():
            la_v0 = "-v0" in ten_luong
            if (ver == "v0") == la_v0 and "v1a" not in ten_luong:
                flows[ten_luong] = {k: luong[k] for k in ("states", "initial", "labeling", "transitions")}
    return {
        "name": ten, "description": mo_ta, "requirements": reqs,
        "state_machine": state_machine(ver),
        "additional_flows": flows,
        "safety_properties": [f for f, _ in THUOC_TINH],
        "mutual_exclusion": {"atom1": LOAI_TRU[0][0], "atom2": LOAI_TRU[0][1]},
        "functions": [_ham_json(h) for h in HAM[ver]],
        "arithmetic_checks": arith,
        "test_cases": [{"id": i, "requirement_ids": [r.strip() for r in rq.split(",")], "title": td,
                        "preconditions": dk, "steps": b, "expected_result": kq, "priority": p}
                       for i, rq, td, dk, b, kq, p in TEST_CASE],
    }


def ghi_excel(ver, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    wb.remove(wb.active)
    dau = Font(bold=True, color="FFFFFF")
    nen = PatternFill("solid", fgColor="1F4E78")
    gop = Alignment(wrap_text=True, vertical="top")

    def sheet(ten, cot, dong, rong):
        ws = wb.create_sheet(ten)
        ws.append(cot)
        for c in ws[1]:
            c.font, c.fill, c.alignment = dau, nen, gop
        for d in dong:
            ws.append(["" if v is None else v for v in d])
        for i, w in enumerate(rong, start=1):
            ws.column_dimensions[chr(64 + i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = gop
        ws.freeze_panes = "A2"

    ten, mo_ta = THONG_TIN[ver]
    sheet("ThongTin", ["Truong", "GiaTri"], [("name", ten), ("description", mo_ta)], [14, 120])
    sheet("YeuCauChucNang", ["ID", "MoTa", "DanhMuc", "MucUuTien"], YEU_CAU_CHUC_NANG, [9, 100, 12, 11])
    sheet("YeuCauPhiChucNang", ["ID", "MoTa", "MucUuTien"], YEU_CAU_PHI_CHUC_NANG, [9, 100, 11])
    sheet("RangBuoc", ["ID", "MoTa", "MucUuTien"], RANG_BUOC, [9, 100, 11])
    sm = state_machine(ver)
    out = {}
    for a, b in sm["transitions"]:
        out.setdefault(a, []).append(b)
    sheet("MoHinhTrangThai", ["TrangThai", "KhoiTao", "Nhan", "ChuyenDenCacTrangThai", "MaTruongHop", "GhiChu"],
          [(s, "x" if s in sm["initial"] else "", ", ".join(lab), ", ".join(out.get(s, [])), ma, gc)
           for s, lab, ma, gc in TRANG_THAI], [22, 8, 45, 55, 11, 70])
    sheet("ThuocTinhAnToan", ["CongThucCTL", "GhiChu"], THUOC_TINH, [24, 90])
    sheet("LoaiTruLanNhau", ["MenhDe1", "MenhDe2"], LOAI_TRU, [18, 18])
    sheet("HamSinhMa", ["Ten", "NgonNgu", "Loai", "ThamSo", "KichThuocBuffer", "Loi", "GiaTriBiMat1", "GiaTriBiMat2", "GhiChu"],
          HAM[ver], [24, 9, 14, 26, 15, 16, 13, 13, 60])
    sheet("KiemTraSoHoc", ["Ten", "BienSoNguyen", "DieuKien", "BieuThucTran", "GhiChu"], SO_HOC[ver], [20, 24, 80, 26, 55])
    sheet("TestCase", ["ID", "YeuCauLienQuan", "TieuDe", "DieuKienTienQuyet", "CacBuoc", "KetQuaMongDoi", "MucUuTien"],
          [(i, rq, td, dk, "\n".join(b), kq, p) for i, rq, td, dk, b, kq, p in TEST_CASE], [7, 18, 32, 40, 55, 60, 10])
    wb.save(path)


def main():
    os.makedirs(os.path.join(HERE, "mo_hinh_con"), exist_ok=True)
    for ver in ("v0", "v1"):
        with open(os.path.join(HERE, f"thanhtoan_donhang_{ver}.json"), "w", encoding="utf-8") as f:
            json.dump(dac_ta_json(ver), f, ensure_ascii=False, indent=2)
    for key, mh in MO_HINH_CON.items():
        with open(os.path.join(HERE, "mo_hinh_con", key + ".json"), "w", encoding="utf-8") as f:
            json.dump(mh, f, ensure_ascii=False, indent=2)
    try:
        for ver in ("v0", "v1"):
            ghi_excel(ver, os.path.join(HERE, f"thanhtoan_donhang_{ver}.xlsx"))
        print("Đã ghi 2 file .xlsx, 2 file .json đặc tả và 6 mô hình con.")
    except ImportError:
        print("Chưa có openpyxl — chỉ ghi được file .json. Cài: pip install openpyxl")


if __name__ == "__main__":
    main()
